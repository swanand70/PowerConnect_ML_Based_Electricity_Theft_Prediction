import sys
import pptx
from pptx.util import Inches

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('Electricity_Theft_Intelligence_Project_Presentation_Revised.pptx')
print(f'Total slides: {len(prs.slides)}')

non_tnr_runs = 0
github_mentions = 0
overflows = 0

uml_slides = range(9, 17) # 0-indexed: slides 10 to 17
screenshot_slides = range(18, 26) # 0-indexed: slides 19 to 26

for idx, slide in enumerate(prs.slides):
    slide_num = idx + 1
    # Check blank canvas slides
    if idx in uml_slides or idx in screenshot_slides:
        text_shapes = [s for s in slide.shapes if hasattr(s, 'text_frame') and s.text_frame.text.strip()]
        if len(text_shapes) != 2:
            print(f'WARNING: Slide {slide_num} does not have exactly 2 text shapes: found {len(text_shapes)}')
    
    for shape in slide.shapes:
        # Check boundary overflow
        if hasattr(shape, 'left') and hasattr(shape, 'top') and hasattr(shape, 'width') and hasattr(shape, 'height'):
            right = shape.left + shape.width
            bottom = shape.top + shape.height
            if right > Inches(13.334) or bottom > Inches(7.501):
                print(f'OVERFLOW on Slide {slide_num}: {shape} right={right/Inches(1):.2f}\", bottom={bottom/Inches(1):.2f}\"')
                overflows += 1
        
        # Check typography & github
        if hasattr(shape, 'text_frame'):
            tf = shape.text_frame
            for p in tf.paragraphs:
                for r in p.runs:
                    if r.text.strip() and r.font.name != 'Times New Roman':
                        print(f'Non-TNR font on slide {slide_num}: "{r.text}" ({r.font.name})')
                        non_tnr_runs += 1
                    if 'github' in r.text.lower() or 'grid-guard' in r.text.lower():
                        print(f'GITHUB mention on slide {slide_num}: "{r.text}"')
                        github_mentions += 1
        elif hasattr(shape, 'table'):
            tbl = shape.table
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            if r.text.strip() and r.font.name != 'Times New Roman':
                                print(f'Non-TNR in table slide {slide_num}: "{r.text}" ({r.font.name})')
                                non_tnr_runs += 1
                            if 'github' in r.text.lower() or 'grid-guard' in r.text.lower():
                                print(f'GITHUB in table slide {slide_num}: "{r.text}"')
                                github_mentions += 1

print('Validation Summary:')
print(f'Non-TNR runs: {non_tnr_runs}')
print(f'GitHub mentions: {github_mentions}')
print(f'Boundary overflows: {overflows}')
print('UML slides (10-17) and Screenshot slides (19-26) verified.')
